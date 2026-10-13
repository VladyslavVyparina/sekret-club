#!/usr/bin/env python3
"""Згенерувати Word-документ з усіма текстами бота Sekret Club."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt, RGBColor

OUT = Path(__file__).resolve().parent.parent / "Sekret_Club_тексти_бота.docx"


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_note(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def add_block(doc, title, when, message, buttons=None, file_ref=None):
    doc.add_paragraph(title, style="List Bullet")
    if when:
        p = doc.add_paragraph()
        p.add_run("Коли показується: ").bold = True
        p.add_run(when)
    if file_ref:
        p = doc.add_paragraph()
        p.add_run("Файл у коді: ").bold = True
        p.add_run(file_ref)
    p = doc.add_paragraph()
    p.add_run("Текст повідомлення:").bold = True
    doc.add_paragraph(message)
    if buttons:
        p = doc.add_paragraph()
        p.add_run("Кнопки:").bold = True
        for b in buttons:
            doc.add_paragraph(b, style="List Bullet 2")
    doc.add_paragraph()


def main():
    doc = Document()
    title = doc.add_heading("Sekret Club — усі тексти бота", 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    add_note(
        doc,
        "Цей документ містить усі тексти, які бачить користувач у Telegram-боті, "
        "а також пов’язані сповіщення про оплату. Редагуйте тексти тут, потім передайте "
        "зміни розробнику для оновлення в коді.\n"
        "Позначки {ЦІНА}, {ДАТА}, {ПОСИЛАННЯ} — динамічні значення (підставляються автоматично).",
    )
    doc.add_page_break()

    # --- 1. Команди ---
    add_heading(doc, "1. Команди в меню Telegram (slash-команди)", 1)
    commands = [
        ("/start", "Почати роботу з ботом"),
        ("/menu", "Головне меню"),
        ("/subscriptions", "Підписка"),
        ("/baza", "База знань"),
        ("/instruction", "Інструкція"),
        ("/oferta", "Договір оферти"),
        ("/audit", "Замовити аудит"),
    ]
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    hdr[0].text = "Команда"
    hdr[1].text = "Опис у меню Telegram"
    for cmd, desc in commands:
        row = t.add_row().cells
        row[0].text = cmd
        row[1].text = desc
    add_note(doc, "Також працюють (без окремого пункту меню): /subscribe, /podpyska, /knowledge, /instructions, /instrukciya, /offer")
    doc.add_paragraph()

    # --- 2. Start ---
    add_heading(doc, "2. Команда /start — вітання та перший екран", 1)

    add_block(
        doc,
        "Trial — унікальне посилання (/start trial)",
        "Новий користувач, ще не користувався trial, не колишній підписник",
        "Вітаю! 👋\n\nЦе бот клубу Sekret Club.\n\nОтримайте 14 днів доступу за 1 грн.",
        buttons=["Отримати доступ за 1 грн"],
        file_ref="bot/bot.py — cmd_start()",
    )

    add_block(
        doc,
        "Новий користувач (звичайний /start)",
        "Без підписки, не колишній підписник",
        "Вітаю! 👋\n\nЦе бот клубу Sekret Club.\n\nОформіть підписку, щоб отримати повний доступ до матеріалів та групи.",
        buttons=["Тариф «Стандарт» — {ЦІНА} грн"],
        file_ref="bot/bot.py — cmd_start()",
    )

    add_block(
        doc,
        "Колишній підписник (/start)",
        "Користувач без активної підписки, статус «колишній підписник»",
        "Раді вас знову бачити!\n\nОформіть підписку, щоб повернутися до клубу.\n\n{ЦІНА} грн/міс — повний доступ.",
        buttons=["Тариф «Стандарт» — {ЦІНА} грн"],
        file_ref="bot/bot.py — cmd_start()",
    )

    add_block(
        doc,
        "Активний підписник (/start)",
        "Є активна підписка — одразу показується головне меню (розділ 3)",
        "—",
        file_ref="bot/bot.py — cmd_start()",
    )

    # --- 3. Main menu ---
    add_heading(doc, "3. Головне меню", 1)

    add_block(
        doc,
        "Заголовок головного меню",
        "Після /start (з підпискою), /menu, кнопка «Назад»",
        "Оберіть пункт меню:",
        buttons=[
            "📚 База знань (лише підписники)",
            "📋 Підписка",
            "📖 Інструкція (лише підписники)",
            "📄 Договір оферти",
            "🎯 Замовити аудит життя (лише підписники)",
            "⬅ Назад (у підменю)",
        ],
        file_ref="bot/bot.py, keyboards.py",
    )

    add_block(
        doc,
        "/menu без підписки",
        "Користувач без активної підписки",
        "Головне меню доступне лише після оформлення підписки.\n\nНатискайте, щоб оформити підписку ⬇️\n/subscriptions",
        file_ref="bot/bot.py — cmd_menu()",
    )

    # --- 4. Subscription ---
    add_heading(doc, "4. Розділ «Підписка» (/subscriptions)", 1)

    add_block(
        doc,
        "Активна підписка — Trial",
        "Користувач на пробному тарифі",
        "📋 Ваш пакет підписки «{НАЗВА_TRIAL} (14 днів)»\n\n"
        "Вартість: 1 грн / 14 днів\n"
        "Строк дії до: {ДАТА}\n"
        "Залишилось днів: {N}\n\n"
        "Дата наступного списання: {ДАТА}\n"
        "(рядок про дату списання не показується, якщо підписку скасовано)",
        buttons=["❌ Скасувати підписку", "⬅ Назад"],
        file_ref="bot/bot.py — _show_subscription()",
    )

    add_block(
        doc,
        "Активна підписка — Стандарт",
        "Користувач на стандартному тарифі",
        "📋 Ваш пакет підписки «{НАЗВА_СТАНДАРТ}»\n\n"
        "Вартість: {ЦІНА} грн / міс.\n"
        "Строк дії до: {ДАТА}\n"
        "Залишилось днів: {N}\n\n"
        "ℹ️ Тариф за 1 грн завершився. Підписку змінено на «{НАЗВА_СТАНДАРТ}» за {ЦІНА} грн/міс.\n"
        "(цей блок — лише якщо раніше був trial)\n\n"
        "Дата наступного списання: {ДАТА}",
        buttons=["❌ Скасувати підписку", "⬅ Назад"],
        file_ref="bot/bot.py — _show_subscription()",
    )

    add_block(
        doc,
        "Без підписки — новий користувач",
        "Розділ Підписка, немає активної підписки",
        "Оформіть підписку.\n\n{ЦІНА} грн/міс — повний доступ до клубу.",
        buttons=["Тариф «Стандарт» — {ЦІНА} грн", "⬅ Назад"],
        file_ref="bot/bot.py — _show_subscription()",
    )

    add_block(
        doc,
        "Без підписки — колишній підписник",
        "Розділ Підписка, статус колишній підписник",
        "Щоб повернутися до клубу — оформіть підписку.\n\n{ЦІНА} грн/міс — повний доступ.",
        buttons=["Тариф «Стандарт» — {ЦІНА} грн", "⬅ Назад"],
        file_ref="bot/bot.py — _show_subscription()",
    )

    add_heading(doc, "4.1. Скасування підписки", 2)
    add_block(
        doc,
        "Підтвердження скасування",
        "Натиснуто «Скасувати підписку»",
        "Ви впевнені, що хочете скасувати підписку?",
        buttons=["✅ Так", "❌ Ні"],
        file_ref="bot/bot.py — cb_sub_cancel()",
    )
    add_block(doc, "Відміна скасування", "Натиснуто «Ні»", "Скасування відмінено.", file_ref="toast-повідомлення")
    add_block(doc, "Підтверджене скасування", "Натиснуто «Так»", "Підписку скасовано. Дякуємо!", file_ref="bot/bot.py — cb_sub_cancel_yes()")
    add_block(doc, "Помилка", "Користувача не знайдено в БД", "Користувача не знайдено.", file_ref="alert")

    # --- 5. Payments ---
    add_heading(doc, "5. Оплата та автоматичні сповіщення", 1)

    add_block(
        doc,
        "Успішна оплата Trial (1 грн)",
        "Перша оплата пробного періоду",
        "✅ Оплата успішна!\n\nІнструкція буде скоро.\n[Тут буде відео-інструкція]",
        buttons=["Перейти в групу"],
        file_ref="web/wayforpay_server.py — process_payment()",
    )

    add_block(
        doc,
        "Успішна оплата / поновлення Стандарт (249 грн)",
        "Автосписання, retry, ручна оплата стандарту",
        "✅ Ваш тариф було оплачено.\n\nТариф «Стандарт» оновлено. Доступ діє до {ДАТА}.",
        file_ref="services/payment_notifications.py — tariff_renewed_message()",
    )

    add_block(
        doc,
        "Нагадування за 1 день до автосписання (Trial)",
        "День 13 після оплати 1 грн — щодня о 10:00",
        "⏰ Нагадування: завтра ({ДАТА}) закінчується пробний період і буде автоматичне списання {ЦІНА} грн за тариф «Стандарт».",
        file_ref="services/payment_notifications.py — charge_tomorrow_message()",
    )

    add_block(
        doc,
        "Нагадування за 1 день до автосписання (Стандарт)",
        "День 29 після попереднього списання — щодня о 10:00",
        "⏰ Нагадування: завтра ({ДАТА}) буде автоматичне списання {ЦІНА} грн за продовження тарифу «Стандарт».",
        file_ref="services/payment_notifications.py — charge_tomorrow_message()",
    )

    add_block(
        doc,
        "Нагадування за 3 дні до кінця підписки",
        "Щодня о 10:00",
        "⏰ Нагадування: ваша підписка закінчується через 3 дні ({ДАТА}). Продовжіть, щоб не втратити доступ до клубу!",
        file_ref="scheduler_tasks.py — remind_subscription_ending()",
    )

    add_block(
        doc,
        "Платіж не пройшов",
        "Відхилена оплата, 5 хв без оплати, retry",
        "Платіж не пройшов, спробуйте ще раз 🔄\nАбо напишіть в підтримку @sekret_manager",
        buttons=["Отримати доступ за 1 грн", "Тариф «Стандарт» — {ЦІНА} грн"],
        file_ref="web/wayforpay_server.py — payment_failed_text()",
    )

    add_block(
        doc,
        "Автосписання після Trial відхилено",
        "Не вдалося списати 249 грн на 14-й день",
        "🚫 Автосписання після Trial відхилено.\n\nТариф: Стандарт\nВартість: {ЦІНА},00 ₴/міс\nНаступна спроба: {ДЕНЬ} {МІСЯЦЬ}",
        file_ref="scheduler_tasks.py — run_trial_to_standard_autocharges()",
    )

    add_block(
        doc,
        "Автосписання Стандарт відхилено",
        "Не вдалося списати на 30-й день циклу",
        "🚫 Автосписання відхилено.\n\nТариф: Стандарт\nВартість: {ЦІНА},00 ₴/міс\nНаступна спроба: {ДЕНЬ} {МІСЯЦЬ}",
        file_ref="scheduler_tasks.py — run_custom_standard_recurring_autocharges()",
    )

    add_block(
        doc,
        "Trial завершився без rec_token",
        "Немає токена для автосписання",
        "⚠️ Пробний період завершився, але автосписання недоступне. Оформіть тариф «Стандарт» у боті, щоб зберегти доступ.",
        file_ref="scheduler_tasks.py",
    )

    add_block(
        doc,
        "Усі спроби списання невдалі",
        "Після 6 невдалих retry — доступ припинено",
        "🚫 Всі спроби списання невдалі. Підписка припинена. Оформіть підписку заново через бота.",
        file_ref="scheduler_tasks.py — run_payment_retries()",
    )

    add_heading(doc, "5.1. Назви товарів на сторінці WayForPay", 2)
    add_block(doc, "Trial", "Сторінка оплати 1 грн", "Підписка на Sekret Club Пробний період", file_ref="web/wayforpay_server.py")
    add_block(doc, "Стандарт", "Сторінка оплати 249 грн", "Підписка на Sekret Club Стандарт / Підписка на Sekret Club", file_ref="web/wayforpay_server.py")

    # --- 6. Group ---
    add_heading(doc, "6. Перехід у групу", 1)
    add_block(
        doc,
        "Повідомлення з посиланням",
        "Після оплати або кнопка «Перейти в групу»",
        "Натисніть кнопку нижче, щоб перейти в групу:\n\n{ПОСИЛАННЯ_НА_ГРУПУ}",
        buttons=["Перейти в групу Sekret Club", "Перейти в групу (після оплати trial)"],
        file_ref="bot/bot.py — cb_join_group()",
    )
    add_block(doc, "Немає підписки", "Спроба перейти без підписки", "Доступ до групи потребує активної підписки.", file_ref="alert")
    add_block(doc, "Немає посилання", "Не налаштовано CLUB_GROUP_LINK", "Помилка. Посилання на групу відсутнє.", file_ref="alert")

    # --- 7. Knowledge base ---
    add_heading(doc, "7. База знань", 1)
    add_block(
        doc,
        "Доступ заборонено",
        "Без підписки",
        "База знань доступна лише підписникам.\n"
        "або: База знань доступна лише підписникам. Оформіть підписку.",
        file_ref="bot/bot.py",
    )
    add_block(
        doc,
        "Головний екран бази знань",
        "Підписник відкрив /baza або кнопку меню",
        "📚 База знань – це всі навчальні матеріали клубу. Оберіть, що вас цікавить:\n"
        "або: 📚 База знань – оберіть категорію:",
        buttons=["📗 Курси", "🎬 Ефіри", "📎 Корисні матеріали", "📚 Книги", "⬅ Назад"],
        file_ref="bot/bot.py, keyboards.py",
    )

    add_heading(doc, "7.1. Курси", 2)
    for text, when in [
        ("Поки що курсів немає.", "Порожній список"),
        ("Оберіть курс:", "Є курси"),
        ("Некоректний курс.", "Помилка callback"),
        ("Курс не знайдено.", "Курс видалено з БД"),
    ]:
        add_block(doc, when, when, text, buttons=["⬅ Назад"] if "немає" in text else None, file_ref="bot/bot.py")

    add_block(
        doc,
        "Кнопка уроку в списку",
        "Шаблон назви кнопки",
        "▶️ Урок №{N}: {скорочена назва уроку}",
        file_ref="keyboards.py — course_lessons()",
    )

    add_heading(doc, "7.2. Ефіри", 2)
    for text, when in [
        ("Поки що ефірів немає.", "Порожній список"),
        ("Оберіть ефір:", "Є ефіри"),
        ("Некоректний ефір.", "Помилка"),
        ("Ефір не знайдено.", "Не знайдено"),
    ]:
        add_block(doc, when, when, text, file_ref="bot/bot.py")

    add_heading(doc, "7.3. Корисні матеріали", 2)
    for text, when in [
        ("Поки що корисних матеріалів немає.", "Порожній список"),
        ("Оберіть матеріал:", "Є матеріали"),
        ("Некоректний матеріал.", "Помилка"),
        ("Матеріал не знайдено.", "Не знайдено"),
    ]:
        add_block(doc, when, when, text, file_ref="bot/bot.py")

    add_heading(doc, "7.4. Книги", 2)
    for text, when in [
        ("Поки що книг немає.", "Порожній список"),
        ("Книги рекомендовані автором:", "Є книги"),
        ("Некоректна книга.", "Помилка"),
        ("Книгу не знайдено.", "Не знайдено"),
    ]:
        add_block(doc, when, when, text, file_ref="bot/bot.py")

    add_heading(doc, "7.5. Шаблон картки контенту (курс/урок/ефір/матеріал/книга)", 2)
    add_block(
        doc,
        "Формат повідомлення",
        "Відображення матеріалу в боті",
        "Курс: «{назва курсу}»\n"
        "(рядок «Курс» — лише для уроків)\n\n"
        "Назва: {назва}\n\n"
        "Опис: {опис}\n"
        "(рядок «Опис» — лише якщо є текст)\n\n"
        "▶️ Відео: {посилання}",
        file_ref="services/content_messages.py — format_entity_message()",
    )

    add_heading(doc, "7.6. Початковий контент бази знань (за замовчуванням)", 2)
    add_note(doc, "Цей контент можна змінити в адмін-панелі. Нижче — тексти з файлу data/knowledge_base.py.")
    courses_seed = [
        'Курс "Вихід на новий рівень"',
        'Курс "Створення реальності"',
        'Курс "Струнке тіло"',
        'Курс "Чиста свідомість"',
        'Курс "Приборкання Ума"',
    ]
    for c in courses_seed:
        doc.add_paragraph(c, style="List Bullet")
    doc.add_paragraph("Уроки курсу «Вихід на новий рівень»:", style="List Bullet")
    lessons = [
        "🎬Лекція №1: Постановка цілей. Як ставити цілі, щоб досягати їх зі 100% гарантією",
        "🎬Лекція №2: Перше «Я» — Душа. Навчитись відчувати душу та виконувати її цілі, плани та задуми",
        '🎬Лекція №3: Друге "Я": УМ. Як приборкати свій ум на шляху до великих досягнень',
        "🎬Лекція №4: Третє \"Я\": Тіло. Як отримувати тіло у високій енергії та стані мобілізації у досягненні цілей",
        "🎬Лекція №5: Контекст: Створення середовища та стану, який сприятиме нові моделі поведінки та мислення",
    ]
    for l in lessons:
        doc.add_paragraph(l, style="List Bullet 2")
    doc.add_paragraph("Ефіри: Ефір №1, Ефір №2, Ефір №3", style="List Bullet")
    doc.add_paragraph("Матеріали: Корисний матеріал 1, Корисний матеріал 2", style="List Bullet")
    doc.add_paragraph("Книги: Книга 1, Книга 2", style="List Bullet")
    doc.add_paragraph()

    # --- 8. Instruction, offer, audit ---
    add_heading(doc, "8. Інструкція, оферта, аудит", 1)
    add_block(doc, "Інструкція — без підписки", "/instruction", "Інструкція доступна лише підписникам.", file_ref="bot/bot.py")
    add_block(
        doc,
        "Інструкція — з посиланням",
        "Підписник, налаштовано INSTRUCTION_DOC_URL",
        "📖 Інструкція до клубу та бота:\n\n{ПОСИЛАННЯ}",
        file_ref="bot/bot.py",
    )
    add_block(
        doc,
        "Інструкція — без посилання",
        "Підписник, URL не налаштовано",
        "📖 Інструкція до клубу та бота:\n\n[Тут може бути відео або посилання]",
        file_ref="bot/bot.py",
    )
    add_block(
        doc,
        "Договір оферти — з посиланням",
        "/oferta, для всіх користувачів",
        "📄 Договір оферти:\n\n{ПОСИЛАННЯ}",
        file_ref="bot/bot.py",
    )
    add_block(
        doc,
        "Договір оферти — без посилання",
        "URL не налаштовано",
        "📄 Договір оферти:\n\n[Посилання на документ]",
        file_ref="bot/bot.py",
    )
    add_block(doc, "Аудит — без підписки", "/audit", "Аудит доступний лише підписникам.", file_ref="bot/bot.py")
    add_block(
        doc,
        "Аудит — основне повідомлення",
        "Підписник",
        "🎯 Замовити аудит життя\n\n"
        "Заповніть заявку на Аудит-сесію — форма нижче (кнопка або посилання).\n\n"
        "{ПОСИЛАННЯ_НА_ФОРМУ}",
        buttons=["📝 Заявка на Аудит-сесію", "⬅ Назад"],
        file_ref="bot/bot.py — AUDIT_MESSAGE_TEXT",
    )

    # --- 9. Admin telegram ---
    add_heading(doc, "9. Адмін-команди в Telegram", 1)
    add_block(doc, "/admin — немає доступу", "Не адмін", "⛔ Немає доступу до адмін-панелі.", file_ref="bot/bot.py")
    add_block(doc, "/admin — посилання", "Адмін", "🔐 Адмін-панель:\n{URL}", file_ref="bot/bot.py")
    add_block(doc, "/admin — помилка URL", "Не налаштовано BASE_URL", "Вкажіть BASE_URL в .env", file_ref="bot/bot.py")
    add_block(
        doc,
        "/admin_clear_me",
        "Скидання тестового акаунта",
        "✅ Скинуто. Ви тепер «не підписник». Напишіть /start для перевірки оплати.",
        file_ref="bot/bot.py",
    )
    add_block(
        doc,
        "/admin_test_sub",
        "Тестова підписка без оплати",
        "✅ Тестова підписка активована на 30 днів. Напишіть /start",
        file_ref="bot/bot.py",
    )

    # --- 10. Web landing ---
    add_heading(doc, "10. Сторінка сайту (не Telegram)", 1)
    add_block(
        doc,
        "Головна сторінка оплати",
        "Відкрили домен без прямого посилання на оплату",
        "Sekret Club\n\nОплата відкривається з бота після вибору тарифу.\n\n"
        "Кнопка: Відкрити бота\n\n"
        "Якщо ви вже маєте посилання на оплату — відкрийте його напряму.",
        file_ref="web/wayforpay_server.py",
    )
    add_block(doc, "Заголовок сторінки оплати", "Редірект на WayForPay", "Оплата", file_ref="HTML title")

    # --- 11. Config names ---
    add_heading(doc, "11. Назви тарифів (налаштовуються)", 1)
    add_note(
        doc,
        "За замовчуванням: Trial — «Trial», Стандарт — «Стандарт». "
        "Можна змінити в адмін-панелі → Тарифи (назви). "
        "Використовуються в текстах підписки та на кнопках оплати.",
    )

    doc.add_paragraph()
    add_note(
        doc,
        "Розсилки з адмін-панелі — довільний текст, який ви вводите самі при відправці. "
        "Контент курсів/уроків/ефірів/книг редагується в адмін-панелі (веб), а не в цьому файлі.",
    )

    doc.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
